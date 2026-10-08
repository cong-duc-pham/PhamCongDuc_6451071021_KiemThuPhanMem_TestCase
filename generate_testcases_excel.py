import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook and sheet
wb = openpyxl.Workbook()

# Sheet 1: Project Overview & Testing Metadata
ws_info = wb.active
ws_info.title = "Thông tin chung"
ws_info.views.sheetView[0].showGridLines = True

# Sheet 2: Test Cases Specification
ws_tc = wb.create_sheet(title="Test Cases")
ws_tc.views.sheetView[0].showGridLines = True

# Style definitions
font_title = Font(name="Arial", size=16, bold=True, color="1A365D")
font_subtitle = Font(name="Arial", size=11, italic=True, color="4A5568")
font_section = Font(name="Arial", size=12, bold=True, color="1A365D")
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
font_body = Font(name="Arial", size=10, color="2D3748")
font_body_bold = Font(name="Arial", size=10, bold=True, color="2D3748")

fill_header = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
fill_zebra = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
fill_module_login = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
fill_module_sso = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid")
fill_module_fp = PatternFill(start_color="EDFDFD", end_color="EDFDFD", fill_type="solid")
fill_module_sec = PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid")
fill_module_ui = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

thin_border_side = Side(border_style="thin", color="CBD5E0")
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ----------------- POPULATE SHEET 1: GENERAL INFORMATION -----------------
ws_info.merge_cells("B2:G2")
ws_info["B2"] = "TÀI LIỆU KỊCH BẢN KIỂM THỬ PHẦN MỀM (TEST CASES SPECIFICATION)"
ws_info["B2"].font = font_title
ws_info["B2"].alignment = Alignment(horizontal="center", vertical="center")

ws_info.merge_cells("B3:G3")
ws_info["B3"] = "Hệ thống: Văn phòng điện tử - Trường Đại học Giao thông Vận tải (UTC)"
ws_info["B3"].font = font_subtitle
ws_info["B3"].alignment = Alignment(horizontal="center", vertical="center")

info_data = [
    ("URL Hệ thống", "https://vanphongdientu.utc.edu.vn/"),
    ("Môn học", "Kiểm thử phần mềm (Software Testing)"),
    ("Sinh viên thực hiện", "Phạm Công Đức"),
    ("Mã sinh viên", "6451071021"),
    ("Trường", "Trường Đại học Giao thông Vận tải (UTC)"),
    ("Phạm vi kiểm thử", "1. Đăng nhập hệ thống (Form Login chuẩn)\n2. Đăng nhập bằng Google SSO (Email UTC)\n3. Quên mật khẩu / Khôi phục tài khoản\n4. Kiểm thử Bảo mật (Security & Validation)\n5. Giao diện người dùng & Khả năng tương thích (UI/UX & Compatibility)"),
    ("Phương pháp áp dụng", "Phân vùng tương đương (EP), Phân tích giá trị biên (BVA), Đoán lỗi (Error Guessing), Kiểm thử bảo mật cơ bản"),
    ("Môi trường kiểm thử", "Trình duyệt: Google Chrome, MS Edge, Firefox\nĐộ phân giải: 1920x1080 (Desktop), 375x812 (Mobile)"),
]

row_idx = 5
for label, val in info_data:
    ws_info.cell(row=row_idx, column=2, value=label).font = font_body_bold
    ws_info.cell(row=row_idx, column=2).alignment = Alignment(horizontal="left", vertical="top")
    ws_info.cell(row=row_idx, column=2).fill = fill_zebra
    ws_info.cell(row=row_idx, column=2).border = border_cell
    
    ws_info.merge_cells(start_row=row_idx, start_column=3, end_row=row_idx, end_column=7)
    val_cell = ws_info.cell(row=row_idx, column=3, value=val)
    val_cell.font = font_body
    val_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for c in range(3, 8):
        ws_info.cell(row=row_idx, column=c).border = border_cell
    row_idx += 1

ws_info.column_dimensions["B"].width = 24
for col_char in ["C", "D", "E", "F", "G"]:
    ws_info.column_dimensions[col_char].width = 18

# ----------------- POPULATE SHEET 2: TEST CASES -----------------
headers = [
    "Test Case ID",
    "Phân hệ (Module)",
    "Tiêu đề ca kiểm thử (Test Scenario)",
    "Điều kiện tiên quyết (Precondition)",
    "Các bước thực hiện (Test Steps)",
    "Dữ liệu kiểm thử (Test Data)",
    "Kết quả mong đợi (Expected Result)",
    "Loại kiểm thử (Test Type)",
    "Mức ưu tiên (Priority)",
    "Trạng thái (Status)",
    "Ghi chú (Notes)"
]

ws_tc.row_dimensions[1].height = 32
for col_idx, h_text in enumerate(headers, start=1):
    cell = ws_tc.cell(row=1, column=col_idx, value=h_text)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_header
    cell.border = border_cell

test_cases = [
    # --- MODULE 1: STANDARD AUTHENTICATION ---
    (
        "TC_LOG_01",
        "1. Đăng nhập chuẩn",
        "Đăng nhập thành công với tài khoản hợp lệ",
        "Người dùng đã có tài khoản đang hoạt động trong hệ thống. Đang ở trang /Login",
        "1. Nhập Username hợp lệ\n2. Nhập Password chính xác\n3. Bấm nút 'Đăng nhập'",
        "Username: 'gv_utc'\nPassword: 'ValidPassword123@'",
        "Đăng nhập thành công, hệ thống điều hướng vào trang chủ Dashboard Văn phòng điện tử và tạo phiên làm việc (Session).",
        "Functional / Positive",
        "High",
        "Untested",
        "Happy Path"
    ),
    (
        "TC_LOG_02",
        "1. Đăng nhập chuẩn",
        "Đăng nhập thất bại khi nhập sai Mật khẩu",
        "Tài khoản người dùng tồn tại. Đang ở trang /Login",
        "1. Nhập Username hợp lệ\n2. Nhập Password không chính xác\n3. Bấm nút 'Đăng nhập'",
        "Username: 'gv_utc'\nPassword: 'WrongPassword999'",
        "Đăng nhập không thành công. Hiển thị thông báo lỗi 'Tên đăng nhập hoặc mật khẩu không chính xác'. Không để lộ chi tiết trường nào bị sai.",
        "Functional / Negative",
        "High",
        "Untested",
        "Kiểm tra bảo mật thông báo lỗi"
    ),
    (
        "TC_LOG_03",
        "1. Đăng nhập chuẩn",
        "Đăng nhập thất bại khi Tên đăng nhập không tồn tại",
        "Đang ở trang /Login",
        "1. Nhập Username không tồn tại trong CSDL\n2. Nhập Password bất kỳ\n3. Bấm nút 'Đăng nhập'",
        "Username: 'user_not_exist_9999'\nPassword: 'Password123'",
        "Hệ thống từ chối đăng nhập. Hiển thị thông báo chung lỗi xác thực tài khoản, không tiết lộ user có tồn tại hay không.",
        "Functional / Negative",
        "High",
        "Untested",
        "Tránh dò quét tài khoản (User Enumeration)"
    ),
    (
        "TC_LOG_04",
        "1. Đăng nhập chuẩn",
        "Để trống cả Tên đăng nhập và Mật khẩu",
        "Đang ở trang /Login",
        "1. Để trống ô Tên đăng nhập\n2. Để trống ô Mật khẩu\n3. Bấm nút 'Đăng nhập'",
        "Username: '' (rỗng)\nPassword: '' (rỗng)",
        "Hệ thống ngăn chặn gửi form hoặc thông báo yêu cầu nhập đầy đủ thông tin Tên đăng nhập và Mật khẩu.",
        "Validation / Negative",
        "Medium",
        "Untested",
        "Validate form client/server"
    ),
    (
        "TC_LOG_05",
        "1. Đăng nhập chuẩn",
        "Để trống ô Tên đăng nhập, chỉ nhập Mật khẩu",
        "Đang ở trang /Login",
        "1. Bỏ trống ô Tên đăng nhập\n2. Nhập Mật khẩu hợp lệ\n3. Bấm nút 'Đăng nhập'",
        "Username: ''\nPassword: 'ValidPassword123@'",
        "Hệ thống báo lỗi yêu cầu nhập Tên đăng nhập. Con trỏ chuột focus vào ô Tên đăng nhập.",
        "Validation / Negative",
        "Medium",
        "Untested",
        ""
    ),
    (
        "TC_LOG_06",
        "1. Đăng nhập chuẩn",
        "Nhập Tên đăng nhập, để trống Mật khẩu",
        "Đang ở trang /Login",
        "1. Nhập Tên đăng nhập hợp lệ\n2. Bỏ trống ô Mật khẩu\n3. Bấm nút 'Đăng nhập'",
        "Username: 'gv_utc'\nPassword: ''",
        "Hệ thống báo lỗi yêu cầu nhập Mật khẩu. Con trỏ chuột focus vào ô Mật khẩu.",
        "Validation / Negative",
        "Medium",
        "Untested",
        ""
    ),
    (
        "TC_LOG_07",
        "1. Đăng nhập chuẩn",
        "Tên đăng nhập có chứa khoảng trắng ở đầu hoặc cuối (Trim space)",
        "Tài khoản 'gv_utc' tồn tại. Đang ở trang /Login",
        "1. Nhập Tên đăng nhập có khoảng trắng: '  gv_utc  '\n2. Nhập Mật khẩu đúng\n3. Bấm 'Đăng nhập'",
        "Username: '  gv_utc  '\nPassword: 'ValidPassword123@'",
        "Hệ thống tự động cắt bỏ (trim) khoảng trắng thừa và cho phép đăng nhập thành công (hoặc báo lỗi rõ ràng nếu quy chuẩn không cho phép).",
        "Boundary / Usability",
        "Low",
        "Untested",
        "Tăng trải nghiệm người dùng"
    ),
    (
        "TC_LOG_08",
        "1. Đăng nhập chuẩn",
        "Kiểm tra tính phân biệt chữ hoa/thường của Mật khẩu (Case Sensitivity)",
        "Tài khoản 'gv_utc' có mật khẩu 'Password123'. Đang ở trang /Login",
        "1. Nhập Tên đăng nhập đúng\n2. Nhập mật khẩu nhưng đổi chữ hoa thành thường ('password123')\n3. Bấm 'Đăng nhập'",
        "Username: 'gv_utc'\nPassword: 'password123'",
        "Đăng nhập thất bại do mật khẩu có phân biệt chữ hoa/chữ thường.",
        "Functional / Security",
        "High",
        "Untested",
        "Quy tắc an toàn mật khẩu"
    ),
    (
        "TC_LOG_09",
        "1. Đăng nhập chuẩn",
        "Kiểm tra chức năng 'Giữ tôi luôn đăng nhập' (Persistent session)",
        "Đang ở trang /Login",
        "1. Nhập đúng Username & Password\n2. Tích chọn checkbox 'Giữ tôi luôn đăng nhập'\n3. Đăng nhập thành công\n4. Đóng trình duyệt và mở lại trang hệ thống",
        "Checkbox persistent = 1",
        "Phiên đăng nhập được duy trì qua Cookie dài hạn (Remember Token), người dùng không cần đăng nhập lại khi mở lại trình duyệt.",
        "Functional",
        "Medium",
        "Untested",
        "Kiểm tra Cookie thời hạn sống"
    ),
    (
        "TC_LOG_10",
        "1. Đăng nhập chuẩn",
        "Kiểm tra phiên làm việc khi KHÔNG chọn 'Giữ tôi luôn đăng nhập'",
        "Đang ở trang /Login",
        "1. Nhập đúng Username & Password\n2. KHÔNG tích checkbox 'Giữ tôi luôn đăng nhập'\n3. Đăng nhập thành công\n4. Đóng toàn bộ trình duyệt và mở lại",
        "Checkbox persistent = uncheck",
        "Session Cookie bị hủy khi đóng trình duyệt. Khi mở lại website, hệ thống bắt buộc chuyển về trang /Login.",
        "Functional",
        "Medium",
        "Untested",
        "Session Cookie tiêu chuẩn"
    ),
    (
        "TC_LOG_11",
        "1. Đăng nhập chuẩn",
        "Ẩn ký tự tại trường Mật khẩu (Password Masking)",
        "Đang ở trang /Login",
        "1. Nhập các ký tự vào ô Mật khẩu\n2. Quan sát hiển thị ký tự trên màn hình",
        "Password: 'Abc@123456'",
        "Ký tự nhập vào hiển thị dưới dạng dấu chấm tròn (•) hoặc dấu hoa thị (*), thuộc tính type='password'.",
        "UI / Security",
        "High",
        "Untested",
        "Tránh lộ mật khẩu nhìn trộm"
    ),
    (
        "TC_LOG_12",
        "1. Đăng nhập chuẩn",
        "Đăng nhập bằng cách nhấn phím Enter",
        "Đang ở trang /Login",
        "1. Nhập Username hợp lệ\n2. Nhập Password hợp lệ\n3. Nhấn phím Enter trên bàn phím (không click chuột)",
        "Phím Enter",
        "Form đăng nhập được submit tự động và đăng nhập thành công vào hệ thống.",
        "Usability",
        "Medium",
        "Untested",
        "Keyboard accessibility"
    ),

    # --- MODULE 2: GOOGLE OAUTH 2.0 FEDERATION (UTC EMAIL) ---
    (
        "TC_SSO_01",
        "2. Google SSO (Email UTC)",
        "Đăng nhập thành công bằng tài khoản Google Mail trường UTC (@utc.edu.vn)",
        "Người dùng có email cán bộ/giảng viên/sinh viên UTC đang hoạt động",
        "1. Tại trang /Login, click nút 'Đăng nhập bằng e-mail UTC'\n2. Chuyển sang giao diện Google OAuth\n3. Chọn/Nhập tài khoản email @utc.edu.vn và mật khẩu Google\n4. Chấp nhận ủy quyền (nếu lần đầu)",
        "Email: 'gv_abc@utc.edu.vn'\nMật khẩu Google hợp lệ",
        "Xác thực Google thành công, callback redirect về vanphongdientu.utc.edu.vn, hệ thống nhận diện đúng user và đăng nhập vào trang chủ.",
        "Functional / SSO",
        "High",
        "Untested",
        "Single Sign-On chính"
    ),
    (
        "TC_SSO_02",
        "2. Google SSO (Email UTC)",
        "Từ chối đăng nhập khi dùng tài khoản Google cá nhân không thuộc tên miền UTC",
        "Đang ở trang đăng nhập Google sau khi bấm 'Đăng nhập bằng e-mail UTC'",
        "1. Bấm nút 'Đăng nhập bằng e-mail UTC'\n2. Đăng nhập bằng tài khoản Gmail thông thường (@gmail.com)\n3. Chờ Google redirect về hệ thống UTC",
        "Email: 'user_canhan@gmail.com'",
        "Hệ thống từ chối đăng nhập và hiển thị thông báo lỗi: 'Vui lòng sử dụng địa chỉ email do Trường ĐH Giao thông Vận tải cấp (@utc.edu.vn)'.",
        "Security / Negative",
        "High",
        "Untested",
        "Kiểm tra Domain Restriction SSO"
    ),
    (
        "TC_SSO_03",
        "2. Google SSO (Email UTC)",
        "Người dùng bấm Hủy (Cancel/Deny) tại màn hình ủy quyền Google",
        "Đang tại màn hình Google Consent",
        "1. Bấm nút 'Đăng nhập bằng e-mail UTC'\n2. Tại màn hình hỏi quyền của Google, bấm nút 'Hủy' (Cancel)",
        "Hành động người dùng từ chối",
        "Hệ thống điều hướng an toàn trở lại trang /Login, không bị lỗi sập trang (Crash / Exception 500) và hiển thị thông báo hủy thao tác.",
        "Exception Handling",
        "Medium",
        "Untested",
        "Xử lý callback error code"
    ),
    (
        "TC_SSO_04",
        "2. Google SSO (Email UTC)",
        "Tài khoản Google @utc.edu.vn hợp lệ nhưng chưa được phân quyền trong hệ thống Văn phòng điện tử",
        "Tài khoản email sinh viên/cán bộ mới chưa khai báo trên hệ thống nội bộ",
        "1. Đăng nhập qua nút Google SSO bằng email @utc.edu.vn hợp lệ nhưng chưa có trong danh sách tài khoản của hệ thống",
        "Email: 'sv_moi@utc.edu.vn'",
        "Hệ thống thông báo tài khoản chưa được kích hoạt hoặc chưa được phân quyền sử dụng hệ thống văn phòng điện tử, hướng dẫn liên hệ QTV.",
        "Functional / Authorization",
        "Medium",
        "Untested",
        "Kiểm tra liên kết account DB"
    ),

    # --- MODULE 3: PASSWORD RECOVERY (/Login/GetPass) ---
    (
        "TC_FP_01",
        "3. Quên mật khẩu",
        "Chuyển hướng đến trang 'Lấy lại mật khẩu' qua liên kết",
        "Đang ở trang /Login",
        "1. Quan sát phần liên kết trợ giúp\n2. Nhấp chuột vào link 'Bạn quên mật khẩu đăng nhập ?'",
        "Link: /Login/GetPass",
        "Hệ thống chuyển hướng thành công đến URL /Login/GetPass, hiển thị tiêu đề 'Lấy lại mật khẩu' cùng form gồm ô Captcha và Email.",
        "Functional / Navigation",
        "High",
        "Untested",
        ""
    ),
    (
        "TC_FP_02",
        "3. Quên mật khẩu",
        "Gửi yêu cầu lấy lại mật khẩu thành công với Email hợp lệ và Captcha đúng",
        "Email đã tồn tại và kích hoạt trong hệ thống. Đang ở /Login/GetPass",
        "1. Nhìn mã ảnh Captcha và nhập chính xác vào ô 'Mã bảo mật'\n2. Nhập địa chỉ Email hợp lệ\n3. Bấm nút 'Cập nhật'",
        "Captcha: đúng mã trên ảnh\nEmail: 'gv_abc@utc.edu.vn'",
        "Hệ thống gửi email hướng dẫn khôi phục mật khẩu hoặc gửi mật khẩu mới về hộp thư. Hiển thị thông báo gửi thành công.",
        "Functional / Positive",
        "High",
        "Untested",
        "Happy Path module Quên MK"
    ),
    (
        "TC_FP_03",
        "3. Quên mật khẩu",
        "Gửi yêu cầu với Email không tồn tại trong hệ thống + Captcha đúng",
        "Đang ở /Login/GetPass",
        "1. Nhập đúng mã Captcha\n2. Nhập Email chưa từng đăng ký trong hệ thống\n3. Bấm 'Cập nhật'",
        "Captcha: chính xác\nEmail: 'notfound12345@utc.edu.vn'",
        "Hệ thống hiển thị thông báo lỗi 'Email không tồn tại trong hệ thống' (hoặc thông báo bảo mật chung).",
        "Functional / Negative",
        "Medium",
        "Untested",
        ""
    ),
    (
        "TC_FP_04",
        "3. Quên mật khẩu",
        "Nhập sai mã bảo mật (Captcha sai) + Email đúng",
        "Đang ở /Login/GetPass",
        "1. Nhập sai mã Captcha (ví dụ: gõ '000000')\n2. Nhập Email hợp lệ\n3. Bấm 'Cập nhật'",
        "Captcha: '000000' (sai)\nEmail: 'gv_abc@utc.edu.vn'",
        "Hệ thống báo lỗi 'Mã bảo mật không chính xác', đồng thời tự động làm mới (refresh) mã captcha mới để chống brute force.",
        "Validation / Security",
        "High",
        "Untested",
        "Bảo vệ chống Bot Spam"
    ),
    (
        "TC_FP_05",
        "3. Quên mật khẩu",
        "Để trống ô Mã bảo mật (Captcha), chỉ nhập Email",
        "Đang ở /Login/GetPass",
        "1. Để trống ô 'Mã bảo mật'\n2. Nhập Email hợp lệ\n3. Bấm 'Cập nhật'",
        "Captcha: ''\nEmail: 'gv_abc@utc.edu.vn'",
        "Hệ thống báo lỗi yêu cầu nhập mã bảo mật, không thực hiện gửi email.",
        "Validation / Negative",
        "Medium",
        "Untested",
        ""
    ),
    (
        "TC_FP_06",
        "3. Quên mật khẩu",
        "Để trống ô Email, chỉ nhập Captcha",
        "Đang ở /Login/GetPass",
        "1. Nhập mã Captcha đúng\n2. Để trống ô 'Địa chỉ Email'\n3. Bấm 'Cập nhật'",
        "Captcha: đúng\nEmail: ''",
        "Hệ thống báo lỗi yêu cầu nhập địa chỉ Email.",
        "Validation / Negative",
        "Medium",
        "Untested",
        ""
    ),
    (
        "TC_FP_07",
        "3. Quên mật khẩu",
        "Nhập Email sai định dạng cú pháp (Invalid Email Format)",
        "Đang ở /Login/GetPass",
        "1. Nhập Captcha đúng\n2. Nhập các dạng email lỗi: 'abc', 'abc@', 'abc@utc', 'abc.edu.vn'\n3. Bấm 'Cập nhật'",
        "Email: 'testinvalidemail'",
        "Hệ thống cảnh báo định dạng Email không hợp lệ (kiểm tra Regex format).",
        "Boundary / Validation",
        "Medium",
        "Untested",
        "Kiểm thử Regex Email"
    ),
    (
        "TC_FP_08",
        "3. Quên mật khẩu",
        "Kiểm tra liên kết 'Trở lại đăng nhập?'",
        "Đang ở /Login/GetPass",
        "1. Nhấp vào link 'Trở lại đăng nhập?' ở phần trợ giúp",
        "Link: /Login",
        "Trình duyệt điều hướng quay trở lại trang đăng nhập chính (/Login).",
        "Navigation",
        "Low",
        "Untested",
        ""
    ),
    (
        "TC_FP_09",
        "3. Quên mật khẩu",
        "Kiểm tra tải lại ảnh Captcha khi F5 hoặc tải lại trang",
        "Đang ở /Login/GetPass",
        "1. Quan sát hình ảnh Captcha hiện tại\n2. Nhấn F5 tải lại trang (hoặc click vào ảnh nếu có hàm reload)",
        "Hành động Reload",
        "Hình ảnh Captcha mới được tạo ra với chuỗi ký tự khác, mã cũ bị hủy hiệu lực.",
        "Functional / Security",
        "Medium",
        "Untested",
        "Chống tái sử dụng Captcha"
    ),

    # --- MODULE 4: SECURITY & NON-FUNCTIONAL VALIDATION ---
    (
        "TC_SEC_01",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra phòng chống tấn công SQL Injection tại ô Tên đăng nhập",
        "Đang ở trang /Login",
        "1. Tại ô Username, nhập chuỗi SQLi: ' OR '1'='1 --\n2. Nhập Password bất kỳ\n3. Bấm 'Đăng nhập'",
        "Username: ' OR '1'='1 --\nPassword: 'test'",
        "Hệ thống xử lý an toàn (sử dụng Parameterized Query / ORM), từ chối đăng nhập, không trả về lỗi cơ sở dữ liệu (Database syntax error).",
        "Security",
        "High",
        "Untested",
        "OWASP Top 10 - Injection"
    ),
    (
        "TC_SEC_02",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra phòng chống tấn công SQL Injection tại ô Mật khẩu",
        "Đang ở trang /Login",
        "1. Nhập Username hợp lệ: 'admin'\n2. Tại ô Password, nhập chuỗi: ' OR '1'='1\n3. Bấm 'Đăng nhập'",
        "Username: 'admin'\nPassword: ' OR '1'='1",
        "Hệ thống từ chối đăng nhập, thông báo sai mật khẩu, không bị bypass xác thực.",
        "Security",
        "High",
        "Untested",
        "OWASP Top 10 - Injection"
    ),
    (
        "TC_SEC_03",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra phòng chống tấn công Cross-Site Scripting (XSS)",
        "Đang ở trang /Login và /Login/GetPass",
        "1. Nhập mã script vào ô Tên đăng nhập và ô Email: <script>alert('XSS')</script>\n2. Bấm submit form",
        "Payload: <script>alert('XSS')</script>",
        "Mã script không được thực thi trên trình duyệt. Dữ liệu được mã hóa HTML entity hoặc bị lọc bỏ an toàn.",
        "Security",
        "High",
        "Untested",
        "OWASP Top 10 - XSS"
    ),
    (
        "TC_SEC_04",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra phòng chống dò mật khẩu liên tục (Brute Force Attack)",
        "Đang ở trang /Login",
        "1. Nhập 1 Username cụ thể và thử nhập sai mật khẩu liên tiếp nhiều lần (5-10 lần trong 1 phút)",
        "5 - 10 lần đăng nhập sai",
        "Hệ thống kích hoạt cơ chế bảo vệ: tạm khóa tài khoản trong 5-15 phút hoặc yêu cầu mã Captcha / giới hạn tần suất yêu cầu (Rate Limiting).",
        "Security",
        "High",
        "Untested",
        "Chống Brute Force"
    ),
    (
        "TC_SEC_05",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra cơ chế chặn truy cập trái phép và tham số redirect ('r=')",
        "Chưa đăng nhập hệ thống",
        "1. Mở tab mới, dán trực tiếp URL trang nội bộ (ví dụ: https://vanphongdientu.utc.edu.vn/)\n2. Quan sát phản hồi của trang",
        "URL nội bộ khi chưa đăng nhập",
        "Hệ thống tự động phát hiện chưa có session và redirect về: https://vanphongdientu.utc.edu.vn/Login?r=... Sau khi đăng nhập thành công sẽ chuyển tiếp về đúng URL ban đầu.",
        "Security / Navigation",
        "High",
        "Untested",
        "Cơ chế Auth Filter & Redirect param"
    ),
    (
        "TC_SEC_06",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra chứng chỉ bảo mật HTTPS / SSL và cờ Cookie",
        "Mở trang https://vanphongdientu.utc.edu.vn/",
        "1. Kiểm tra biểu tượng ổ khóa HTTPS trên thanh địa chỉ\n2. Mở DevTools > Application > Cookies, kiểm tra cờ thuộc tính của Cookie phiên",
        "Giao thức HTTPS",
        "Website sử dụng kết nối HTTPS hợp lệ, không có cảnh báo 'Not Secure'. Session cookie có gắn cờ HttpOnly (chống trộm cookie qua JS) và Secure flag.",
        "Security",
        "High",
        "Untested",
        "Bảo mật truyền thông mạng"
    ),
    (
        "TC_SEC_07",
        "4. Bảo mật & Phi chức năng",
        "Kiểm tra nút Back của trình duyệt sau khi Đăng xuất (Browser Back Button)",
        "Đã đăng nhập thành công",
        "1. Thực hiện Đăng xuất khỏi hệ thống\n2. Bấm nút Back trên trình duyệt để quay lại trang nội bộ trước đó",
        "Hành động Back trình duyệt",
        "Hệ thống không cho phép xem dữ liệu cũ từ cache trình duyệt, tự động redirect về trang /Login và yêu cầu đăng nhập lại.",
        "Security / Usability",
        "Medium",
        "Untested",
        "Cache-Control: no-cache, no-store"
    ),

    # --- MODULE 5: UI & COMPATIBILITY TESTING ---
    (
        "TC_UI_01",
        "5. Giao diện & Tương thích",
        "Kiểm tra hiển thị đầy đủ các thành phần giao diện trang Login",
        "Mở trang https://vanphongdientu.utc.edu.vn/Login",
        "1. Quan sát banner bên trái (hình ảnh flight.png, slogan 'Không chỉ là một giải pháp quản lý - Làm việc mọi lúc mọi nơi')\n2. Quan sát form bên phải\n3. Quan sát chân trang (Footer)",
        "Giao diện chuẩn",
        "Tất cả hình ảnh, icon, chữ tiếng Việt có dấu hiển thị sắc nét, không bị vỡ font (UTF-8 chuẩn), không lệch layout.",
        "UI",
        "Medium",
        "Untested",
        "Kiểm tra Font & Bố cục"
    ),
    (
        "TC_UI_02",
        "5. Giao diện & Tương thích",
        "Kiểm tra liên kết chân trang 'Trung tâm trợ giúp'",
        "Đang ở trang /Login hoặc /Login/GetPass",
        "1. Cuộn xuống chân trang\n2. Nhấp vào liên kết 'Trung tâm trợ giúp'",
        "Link: http://hotrokythuat.utc.edu.vn",
        "Mở tab mới (target='_blank') điều hướng chính xác đến trang cổng Hỗ trợ kỹ thuật UTC.",
        "UI / Navigation",
        "Low",
        "Untested",
        ""
    ),
    (
        "TC_UI_03",
        "5. Giao diện & Tương thích",
        "Kiểm tra liên kết chân trang 'Ý kiến phản hồi'",
        "Đang ở trang /Login hoặc /Login/GetPass",
        "1. Cuộn xuống chân trang\n2. Nhấp vào liên kết 'Ý kiến phản hồi'",
        "Link: mailto:hotrokythuat@utc.edu.vn",
        "Kích hoạt ứng dụng email mặc định (Outlook/Mail) với địa chỉ người nhận điền sẵn là hotrokythuat@utc.edu.vn.",
        "UI / Usability",
        "Low",
        "Untested",
        ""
    ),
    (
        "TC_UI_04",
        "5. Giao diện & Tương thích",
        "Kiểm tra độ tương thích trên các trình duyệt phổ biến (Cross-browser)",
        "Cài đặt Google Chrome, Microsoft Edge, Mozilla Firefox",
        "1. Mở trang web lần lượt trên Chrome, Edge, Firefox\n2. Thử nghiệm các thao tác nhập liệu và click nút",
        "Các trình duyệt máy tính",
        "Website hoạt động ổn định, hiển thị đồng nhất về màu sắc, kích thước nút và font chữ trên tất cả các trình duyệt.",
        "Compatibility",
        "Medium",
        "Untested",
        "Kiểm thử tương thích trình duyệt"
    ),
    (
        "TC_UI_05",
        "5. Giao diện & Tương thích",
        "Kiểm tra độ co giãn hiển thị trên thiết bị di động / màn hình nhỏ (Responsive)",
        "Sử dụng công cụ Device Mode (F12) trên Chrome hoặc điện thoại thật",
        "1. Điều chỉnh kích thước màn hình về 375x812 (iPhone/Android)\n2. Kiểm tra các trường nhập liệu và nút bấm",
        "Kích thước Mobile/Tablet",
        "Form đăng nhập không bị tràn ngang màn hình, các ô input và nút bấm hiển thị vừa vặn, người dùng bấm chạm dễ dàng.",
        "Responsive / Usability",
        "Medium",
        "Untested",
        "Mobile Viewport"
    )
]

for row_i, tc in enumerate(test_cases, start=2):
    ws_tc.row_dimensions[row_i].height = 54
    module_name = tc[1]
    
    # Decide module background tint
    if "1. Đăng nhập" in module_name:
        row_fill = fill_module_login if row_i % 2 == 0 else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    elif "2. Google SSO" in module_name:
        row_fill = fill_module_sso if row_i % 2 == 0 else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    elif "3. Quên mật khẩu" in module_name:
        row_fill = fill_module_fp if row_i % 2 == 0 else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    elif "4. Bảo mật" in module_name:
        row_fill = fill_module_sec if row_i % 2 == 0 else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    else:
        row_fill = fill_module_ui if row_i % 2 == 0 else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for col_i, val in enumerate(tc, start=1):
        c = ws_tc.cell(row=row_i, column=col_i, value=val)
        c.font = font_body
        c.border = border_cell
        c.fill = row_fill
        
        # Alignment specific rules
        if col_i in [1, 9, 10]:  # ID, Priority, Status
            c.alignment = align_center
            if col_i == 1:
                c.font = font_body_bold
        elif col_i in [2, 8]:   # Module, Test Type
            c.alignment = align_center
        else:
            c.alignment = align_left

# Set column widths for Test Cases sheet
col_widths = {
    "A": 14, # ID
    "B": 22, # Module
    "C": 32, # Test Scenario
    "D": 26, # Precondition
    "E": 36, # Test Steps
    "F": 28, # Test Data
    "G": 38, # Expected Result
    "H": 18, # Test Type
    "I": 12, # Priority
    "J": 12, # Status
    "K": 24  # Notes
}

for col_letter, width in col_widths.items():
    ws_tc.column_dimensions[col_letter].width = width

# Freeze pane at header row
ws_tc.freeze_panes = "A2"

output_path = r"c:\CongDuc\Ki_I_nam_IV\KiemThuPhanMem\PhamCongDuc_6451071021\TestCases_VanPhongDienTu_UTC.xlsx"
wb.save(output_path)
print("Saved Excel successfully to:", output_path)
